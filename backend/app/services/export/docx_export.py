import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor

from app.schemas.resume import ResumeContent
from app.services.export.layout import EntryBlock, contact_line
from app.services.export.sections import ResumeSection, build_sections

BODY_FONT = "Calibri"
BODY_SIZE = Pt(10.5)
NAME_SIZE = Pt(20)
HEADING_SIZE = Pt(11.5)
META_SIZE = Pt(9.5)
PAGE_MARGIN = Inches(0.7)
MUTED = RGBColor(0x55, 0x5B, 0x66)
ACCENT = RGBColor(0x3E, 0x47, 0xA8)
SECTION_SPACE_BEFORE = Pt(10)
ENTRY_SPACE_BEFORE = Pt(6)
BULLET_STYLE = "List Bullet"


def _configure_document(document: DocxDocument) -> None:
    for section in document.sections:
        section.left_margin = section.right_margin = PAGE_MARGIN
        section.top_margin = section.bottom_margin = PAGE_MARGIN
    normal = document.styles["Normal"]
    font = getattr(normal, "font", None)
    if font is not None:
        font.name = BODY_FONT
        font.size = BODY_SIZE


def _add_header(document: DocxDocument, resume: ResumeContent) -> None:
    name_run = document.add_paragraph().add_run(resume.name)
    name_run.bold = True
    name_run.font.size = NAME_SIZE
    contact = contact_line(resume.contact)
    if contact:
        contact_run = document.add_paragraph().add_run(contact)
        contact_run.font.size = META_SIZE
        contact_run.font.color.rgb = MUTED


def _add_section_heading(document: DocxDocument, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = SECTION_SPACE_BEFORE
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = HEADING_SIZE
    run.font.color.rgb = ACCENT


def _add_entry(document: DocxDocument, entry: EntryBlock) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_before = ENTRY_SPACE_BEFORE
    title.paragraph_format.keep_with_next = True
    title.add_run(entry.heading).bold = True
    if entry.subheading:
        title.add_run(f"  |  {entry.subheading}")
    if entry.meta:
        meta_run = document.add_paragraph().add_run(entry.meta)
        meta_run.italic = True
        meta_run.font.size = META_SIZE
        meta_run.font.color.rgb = MUTED
    for line in entry.lines:
        document.add_paragraph(line, style=BULLET_STYLE)


def _add_section(document: DocxDocument, section: ResumeSection) -> None:
    _add_section_heading(document, section.title)
    if section.paragraph:
        document.add_paragraph(section.paragraph)
    if section.inline_list:
        document.add_paragraph(", ".join(section.inline_list))
    for entry in section.entries:
        _add_entry(document, entry)


def build_docx(resume: ResumeContent) -> bytes:
    document = Document()
    _configure_document(document)
    _add_header(document, resume)
    for section in build_sections(resume):
        _add_section(document, section)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
