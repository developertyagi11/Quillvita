"""Builders for realistic test inputs and canned LLM output."""

import io
import json
from collections.abc import Sequence

from docx import Document
from fpdf import FPDF
from fpdf.enums import XPos, YPos

RESUME_LINES: tuple[str, ...] = (
    "Jane Doe",
    "jane@example.com | +1 555 0100 | Berlin",
    "SUMMARY",
    "Backend engineer with 5 years of experience building Python services.",
    "EXPERIENCE",
    "Senior Engineer | Acme Corp | 2021 - Present",
    "- Built REST APIs with FastAPI serving 2M requests per day",
    "- Reduced infrastructure cost by 30% by migrating to containers",
)

JOB_DESCRIPTION = (
    "We are hiring a Senior Backend Engineer to design and scale Python APIs. "
    "You will work with FastAPI, PostgreSQL, Docker and Kubernetes, own services end to end, "
    "and mentor engineers. Experience with cloud cost optimisation is a plus."
)


def make_docx(paragraphs: Sequence[str], table_rows: Sequence[Sequence[str]] = ()) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(lines: Sequence[str]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.cell(0, 8, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if not lines:
        pdf.rect(20, 20, 100, 60, style="F")
    return bytes(pdf.output())


def resume_payload(summary: str, skills: Sequence[str]) -> dict[str, object]:
    return {
        "name": "Jane Doe",
        "contact": {"email": "jane@example.com", "phone": "", "location": "Berlin", "links": []},
        "summary": summary,
        "skills": list(skills),
        "experience": [
            {
                "title": "Senior Engineer",
                "company": "Acme Corp",
                "location": "",
                "start_date": "2021",
                "end_date": "Present",
                "bullets": ["Built REST APIs with FastAPI serving 2M requests per day"],
            }
        ],
        "education": [],
        "projects": [],
    }


def customization_payload(tailored_skills: Sequence[str] = ("FastAPI", "Python")) -> str:
    return json.dumps(
        {
            "original": resume_payload("Backend engineer.", ["Python", "FastAPI"]),
            "tailored": resume_payload("Senior backend engineer scaling APIs.", tailored_skills),
            "matched_skills": ["FastAPI", "Python", "FastAPI"],
            "suggested_keywords": [{"keyword": "Kubernetes", "reason": "Required in the JD."}],
            "scores": {
                "match_score": 78,
                "match_justification": "Strong API background; no Kubernetes.",
                "ats_score": 82,
                "ats_tips": ["Add a dedicated Tools line."],
            },
            "changes": [
                {
                    "section": "summary",
                    "before": "Backend engineer.",
                    "after": "Senior backend engineer scaling APIs.",
                    "reason": "Mirrors the JD title.",
                },
                {
                    "section": "skills",
                    "before": "Python, FastAPI",
                    "after": "FastAPI, Python",
                    "reason": "FastAPI is listed first in the JD.",
                },
            ],
        }
    )
